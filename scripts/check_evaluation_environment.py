#!/usr/bin/env python3
"""Smoke-check PDF/PPTX dependencies using temporary synthetic inputs only."""

from __future__ import annotations

import importlib.metadata
import json
import platform
import subprocess
import sys
import tempfile
from pathlib import Path

import fitz
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from check_deck_quality import pdf_metrics, pptx_geometry


def main() -> int:
    project_dir = Path(__file__).resolve().parent.parent
    with tempfile.TemporaryDirectory(prefix="survey-env-smoke-") as temp_dir:
        root = Path(temp_dir)
        pdf_path = root / "synthetic.pdf"
        image_path = root / "cutout.png"
        with fitz.open() as document:
            page = document.new_page(width=300, height=200)
            page.insert_text((25, 35), "Environment smoke test")
            page.draw_rect(fitz.Rect(40, 60, 240, 160), color=(0, 0, 1))
            document.save(pdf_path)

        subprocess.run(
            [
                sys.executable, str(project_dir / "scripts/extract_source_figure.py"),
                "pdf", "--source", str(pdf_path), "--page", "1",
                "--bbox", "40,60,240,160", "--dpi", "72",
                "--output", str(image_path),
            ],
            check=True, capture_output=True, text=True,
        )
        with Image.open(image_path) as image:
            if image.size != (200, 100):
                raise RuntimeError(f"Unexpected crop dimensions: {image.size}")
            image.verify()
        manifest = json.loads(image_path.with_suffix(".json").read_text())
        if manifest["page_count"] != 1 or manifest["output_size_px"] != [200, 100]:
            raise RuntimeError("PDF extraction provenance smoke check failed")

        pdf_observation = pdf_metrics(pdf_path, [])
        if not pdf_observation["available"]:
            raise RuntimeError(f"PDF observer unavailable: {pdf_observation}")

        pptx_path = root / "synthetic.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_picture(str(image_path), Inches(1), Inches(1))
        presentation.save(pptx_path)
        pptx_observation = pptx_geometry(pptx_path)
        if not pptx_observation["available"]:
            raise RuntimeError(f"PPTX observer unavailable: {pptx_observation}")

    print(json.dumps({
        "status": "passed",
        "python": platform.python_version(),
        "platform": platform.platform(),
        "dependencies": {
            name: importlib.metadata.version(name)
            for name in ("PyMuPDF", "Pillow", "python-pptx", "lxml", "typing_extensions", "XlsxWriter")
        },
        "checks": ["synthetic PDF create/open/render/crop/provenance", "PDF observer", "PPTX create/observer"],
        "historical_evaluations_run": False,
    }, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
