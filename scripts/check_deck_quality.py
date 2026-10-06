#!/usr/bin/env python3
"""Observe a survey deck and report deterministic release defects.

The manifest is an authoring contract, not implementation evidence. Run the
structured evidence-visual preflight before this observer. This checker
compares the manifest with actual PPTX/PDF objects and reader-facing text; it
does not replace the preflight's route/provenance joins. It hard-fails
deterministic defects and reports semantic questions for review.
Use --min-observed-visuals and --min-body-pt only when the task or selected
presentation profile has justified those thresholds.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
import zipfile
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from xml.etree import ElementTree as ET

from check_review_record import check_review_file
from check_storyline_plot import _load_payload, check_plot_file, plot_rows as plot_rows_from_payload
from manifest_visuals import deck_visual_rows, has_nested_visuals


ALLOWED_VISUALS = {
    "source_figure", "reconstructed_figure", "result_table", "result_plot",
    "method_diagram", "comparison", "equation", "boundary_map",
    "photo_or_screenshot", "none",
}
NON_EVIDENCE_CLASSES = {"cover", "orientation", "transition", "reference"}
EVIDENCE_CLASSES = {"evidence", "method", "boundary", "synthesis"}
SOURCE_VISUAL_ROUTES = {
    "source_figure", "faithful_reconstruction", "new_explanatory_visual",
    "no_useful_candidate", "inaccessible_source", "not_applicable",
}
RECONSTRUCTION_ROUTES = {"faithful_reconstruction", "new_explanatory_visual"}
NON_USE_REASON_CODES = {
    "layout_unreadable",
    "no_source_figure",
    "cross_source_comparison",
    "review_synthesis",
}
VISUAL_REVIEW_STATUSES = {"pending", "passed", "blocked", "not_required", "not_applicable"}
EVIDENCE_CONTRACT_KINDS = {"table", "diagram", "chart", "source_figure"}
RECONSTRUCTION_ATTRIBUTION_PATTERNS = (
    re.compile(r"reconstruct(?:ed|ion).*survey|survey.*reconstruct", re.IGNORECASE),
    re.compile(r"surveys+synthesis", re.IGNORECASE),
    re.compile(r"サーベイ.*(?:再構成|合成)|(?:再構成|合成).*サーベイ"),
    re.compile(r"レビュー側.*(?:再構成|合成)|(?:再構成|合成).*レビュー側"),
)
INTERNAL_TOKENS = (
    "SOURCE-REPORTED", "SOURCE-INFORMED", "MANAGER SYNTHESIS", "WBS",
    "RACI", "Stage-Gate", "manager/worker", "reviewer ID", "source_id",
    "claim_id", "evidence_id", "production instruction", "source_cutout_asset_ids",
    "source_cutout_manifest_paths", "source_scan_record_paths",
    "visual_comparison_packet_path", "visual_review_status", "decision_reason_code",
    "source_visual_route", "derived_asset_id", "derived_asset_path", "attribution_note",
    "source_ids", "claim_ids", "rights_basis", "source_cutout_created_at",
    "reconstruction_created_at", "asset_role", "review_only", "source_sha256",
    "output_sha256", "asset_id", "source_hash_basis", "source_snapshot_path",
    "local_image_path", "local_source_path",
)
DEFAULT_ACCENTS = {"BF0000", "C00000", "DC2626", "B91C1C"}
PDF_BACKGROUND_AREA_RATIO = 0.98
PDF_TINY_DRAWING_AREA_RATIO = 0.005
FORBIDDEN_LINE_START = set("、。，．：；？！)]}〉》」』】〕）］≫〗〙〛")
FORBIDDEN_LINE_END = set("([<{〈《「『【〔（［﹁﹃")
SLIDE_ID_PATTERN = re.compile(r"^(?:\d+|s\d+|slide[-_ ]?\d+)$", re.IGNORECASE)
CITATION_PATTERNS = (
    re.compile(r"(?:https?://|www\.)\S+", re.IGNORECASE),
    re.compile(r"\b(?:doi:\s*)?10\.\d{4,9}/\S+", re.IGNORECASE),
    # Common slide-deck shorthand such as ``J. Smith+`` is a stronger
    # source-credit signal than a venue/event token followed by a year.
    re.compile(r"\b(?:[A-Z]\.\s+){1,4}[\w'’\-]+(?:\+|\s+et\s+al\.)", re.UNICODE),
    # Require a separator before the year.  This intentionally avoids
    # matching tokens such as MIRU2026, CVPR2024, or SIGGRAPH1999.
    re.compile(r"\b[A-Z][A-Za-z'’\-]+(?:\s+et\s+al\.)?[\s,\(\[\{]+(?:19|20)\d{2}[\)\]\}]?(?=\s|$|[.,;:])"),
    re.compile(r"[\u3040-\u30ff\u3400-\u9fff]{2,15}\s*[（(][12]\d{3}[）)]"),
)
NUMERIC_PATTERNS = (
    re.compile(
        r"(?<![\w])[-+]?\d+(?:[.,]\d+)?\s*(?:%|％|percent|percentage|‰|°C|°F|"
        r"ms|msec|s|sec|seconds?|min|minutes?|h|hours?|kg|mg|g|km|cm|mm|m|"
        r"px|pixels?|Hz|kHz|MHz|GB|MB|dB|USD|points?)(?![A-Za-z])",
        re.IGNORECASE,
    ),
    re.compile(r"\b(?:n|N|p|r|d|f|t|CI|M|SD|SE|mean|median|sample|participants?|accuracy|rate|score|effect)\s*[=:]\s*[-+]?\d+(?:[.,]\d+)?", re.IGNORECASE),
    re.compile(
        r"(?<![\w])\d+(?:[.,]\d+)?\s+(?:stud(?:y|ies)|papers?|participants?|people|users?|"
        r"items?|trials?|sources?|works?|reports?|cases?|pages?|models?|methods?|conditions?|"
        r"groups?|classes?|observations?|samples?)(?![\w])",
        re.IGNORECASE,
    ),
    re.compile(r"(?<![\w])\d+\s*/\s*\d+(?![\w])"),
    re.compile(r"\b(?:sample(?:\s+size)?|number|total|count)\s*(?:of|=|:)\s*\d+(?:[.,]\d+)?\b", re.IGNORECASE),
)


def lname(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def natural_key(name: str):
    return [int(x) if x.isdigit() else x for x in re.split(r"(\d+)", name)]


def citation_match(text: str) -> str | None:
    for pattern in CITATION_PATTERNS:
        match = pattern.search(text)
        if match:
            return match.group(0).rstrip(".,;:)]}）")
    return None


def pdf_visual_metrics(page):
    """Count conservative PDF visual objects without treating page chrome as evidence."""
    page_rect = page.rect
    page_area = page_rect.width * page_rect.height
    raw_drawings = page.get_drawings()
    meaningful_drawings = []
    ignored_background = 0
    ignored_tiny = 0
    for drawing in raw_drawings:
        rect = drawing.get("rect")
        if rect is None or rect.width <= 0 or rect.height <= 0 or page_area <= 0:
            ignored_tiny += 1
            continue
        area_ratio = (rect.width * rect.height) / page_area
        covers_page = (
            rect.x0 <= page_rect.x0 + 1
            and rect.y0 <= page_rect.y0 + 1
            and rect.x1 >= page_rect.x1 - 1
            and rect.y1 >= page_rect.y1 - 1
        )
        # Exported slide decks commonly contain a full-page white rectangle as
        # the background. It is an object, but not evidence-bearing content.
        if covers_page and area_ratio >= PDF_BACKGROUND_AREA_RATIO:
            ignored_background += 1
        # Thin rules, page chrome, and other sub-0.5% primitives should not
        # make an evidence row pass by themselves. Larger vector diagrams
        # remain observable; semantic adequacy is still review-only.
        elif area_ratio < PDF_TINY_DRAWING_AREA_RATIO:
            ignored_tiny += 1
        else:
            meaningful_drawings.append(drawing)
    image_count = len(page.get_images(full=True))
    return {
        "images": image_count,
        "drawings": len(raw_drawings),
        "observed_vector_drawings": len(meaningful_drawings),
        "ignored_background_drawings": ignored_background,
        "ignored_tiny_drawings": ignored_tiny,
        "observed_visuals": image_count + len(meaningful_drawings),
        "has_observed_visual": image_count > 0 or bool(meaningful_drawings),
    }


def numeric_matches(text: str) -> list[str]:
    matches = []
    for pattern in NUMERIC_PATTERNS:
        matches.extend(match.group(0).strip() for match in pattern.finditer(text))
    return sorted(set(matches))


def is_provenance_only(text: str) -> bool:
    remainder = text
    for pattern in CITATION_PATTERNS:
        remainder = pattern.sub(" ", remainder)
    remainder = re.sub(r"\b(?:doi|source|fig(?:ure)?)\b", " ", remainder, flags=re.IGNORECASE)
    remainder = re.sub(r"[^\w\u3040-\u30ff\u3400-\u9fff]+", "", remainder)
    return len(remainder) < 12


def slide_number(slide_id: str) -> int | None:
    if not SLIDE_ID_PATTERN.fullmatch(slide_id or ""):
        return None
    match = re.search(r"(\d+)$", slide_id)
    return int(match.group(1)) if match else None


def pipe_values(value: str) -> list[str]:
    return [item.strip() for item in (value or "").split("|") if item.strip()]


def valid_timestamp(value: str) -> bool:
    if not value.strip():
        return False
    try:
        datetime.fromisoformat(value.strip().replace("Z", "+00:00"))
    except ValueError:
        return False
    return True


def timestamp_value(value: str) -> datetime | None:
    if not valid_timestamp(value):
        return None
    parsed = datetime.fromisoformat(value.strip().replace("Z", "+00:00"))
    return parsed if parsed.tzinfo is not None else parsed.replace(tzinfo=timezone.utc)


def compact_text(value: str) -> str:
    return re.sub(r"\s+", "", value or "").casefold()


def attribution_observed(page_text: str, attribution_note: str, *, allow_generic: bool = True) -> bool:
    page_compact = compact_text(page_text)
    note_compact = compact_text(attribution_note)
    if note_compact and note_compact in page_compact:
        return True
    return allow_generic and any(pattern.search(page_text or "") for pattern in RECONSTRUCTION_ATTRIBUTION_PATTERNS)


def audit_paths_exist(raw_paths: str, audit_root: Path | None) -> tuple[list[str], list[str]]:
    """Return missing and uncheckable paths without treating URLs as local files."""
    if not audit_root:
        return [], []
    missing = []
    uncheckable = []
    for raw in pipe_values(raw_paths):
        if re.match(r"^[a-z][a-z0-9+.-]*://", raw, re.IGNORECASE):
            uncheckable.append(raw)
            continue
        candidate = Path(raw)
        if not candidate.is_absolute():
            candidate = audit_root / candidate
        if not candidate.is_file():
            missing.append(str(candidate))
    return missing, uncheckable


def audit_paths_outside_root(raw_paths: str, audit_root: Path | None) -> list[str]:
    if audit_root is None:
        return []
    outside = []
    for raw in pipe_values(raw_paths):
        path = resolve_audit_path(raw, audit_root)
        if path is not None and not path_within_audit_root(path, audit_root):
            outside.append(str(path))
    return outside


def resolve_audit_path(raw_path: str, audit_root: Path | None) -> Path | None:
    """Resolve a local audit path; URLs are intentionally not materialized."""
    if re.match(r"^[a-z][a-z0-9+.-]*://", raw_path, re.IGNORECASE):
        return None
    candidate = Path(raw_path)
    if not candidate.is_absolute():
        if audit_root is None:
            return None
        candidate = audit_root / candidate
    return candidate


def path_within_audit_root(path: Path, audit_root: Path | None) -> bool:
    if audit_root is None:
        return True
    try:
        path.resolve().relative_to(audit_root.resolve())
    except ValueError:
        return False
    return True


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def truthy(value) -> bool:
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in {"true", "yes", "1"}


def record_values(record: dict, key: str) -> list[str]:
    value = record.get(key, [])
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]
    return pipe_values(str(value))


def load_json_records(raw_paths: str, audit_root: Path | None, label: str, strict: bool):
    """Load local JSON audit records and return (records, errors, warnings)."""
    records = []
    errors = []
    warnings = []
    for raw in pipe_values(raw_paths):
        if audit_root is None:
            errors.append(f"{label} requires --audit-root before referenced records can be read")
            continue
        path = resolve_audit_path(raw, audit_root)
        if path is None:
            message = f"{label} path is not locally inspectable: {raw}"
            errors.append(message)
            continue
        if not path_within_audit_root(path, audit_root):
            message = f"{label} path is outside audit-root: {path}"
            (errors if strict else warnings).append(message)
            continue
        if not path.is_file():
            message = f"{label} path does not exist: {path}"
            (errors if strict else warnings).append(message)
            continue
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, UnicodeError, json.JSONDecodeError) as exc:
            errors.append(f"{label} is not valid JSON ({path}): {exc}")
            continue
        if isinstance(payload, list):
            if len(payload) != 1 or not isinstance(payload[0], dict):
                errors.append(f"{label} must contain exactly one JSON object per visual ({path})")
                continue
            payload = payload[0]
        if not isinstance(payload, dict):
            errors.append(f"{label} must contain a JSON object ({path})")
            continue
        records.append((raw, path, payload))
    return records, errors, warnings


def check_source_cutout_records(row, source_cutout_ids, source_cutout_paths, audit_root, strict):
    """Check source-cutout sidecars and their joins to the manifest row."""
    errors = []
    warnings = []
    records, load_errors, load_warnings = load_json_records(
        row.get("source_cutout_manifest_paths", ""),
        audit_root,
        f"slide {row.get('slide_id', '?')} source-cutout manifest",
        strict,
    )
    errors.extend(load_errors)
    warnings.extend(load_warnings)
    if len(source_cutout_ids) != len(source_cutout_paths):
        errors.append(
            f"slide {row.get('slide_id', '?')}: source_cutout_asset_ids and source_cutout_manifest_paths must have equal counts"
        )
    record_ids = []
    row_figure_ids = set(pipe_values(row.get("figure_id", "")))
    row_source_ids = set(pipe_values(row.get("source_ids", "")))
    row_claim_ids = set(pipe_values(row.get("claim_ids", "")))
    expected_visual_id = str(row.get("visual_id", "")).strip()
    expected_generation_id = str(row.get("generation_id", "")).strip()
    for _raw, path, record in records:
        record_visual_id = str(record.get("visual_id", record.get("visualId", ""))).strip()
        if not expected_visual_id:
            errors.append(f"slide {row.get('slide_id', '?')}: source-cutout validation requires manifest visual_id")
        if not record_visual_id:
            errors.append(f"slide {row.get('slide_id', '?')}: source-cutout sidecar {path} lacks visual_id")
        elif expected_visual_id and record_visual_id != expected_visual_id:
            errors.append(f"slide {row.get('slide_id', '?')}: source-cutout visual_id does not match the manifest")
        if expected_generation_id:
            record_generation_id = str(record.get("generation_id", record.get("generationId", ""))).strip()
            if not record_generation_id:
                errors.append(f"slide {row.get('slide_id', '?')}: source-cutout sidecar {path} lacks generation_id")
            elif record_generation_id != expected_generation_id:
                errors.append(f"slide {row.get('slide_id', '?')}: source-cutout generation_id does not match the manifest")
        elif str(record.get("generation_id", record.get("generationId", ""))).strip():
            errors.append(f"slide {row.get('slide_id', '?')}: source-cutout generation_id is not joined to the manifest")
        asset_id = str(record.get("asset_id", "")).strip()
        record_ids.append(asset_id)
        if not asset_id or asset_id not in source_cutout_ids:
            errors.append(f"slide {row.get('slide_id', '?')}: source-cutout sidecar {path} has an ID not listed in the manifest")
        if record.get("asset_role") != "source_cutout":
            errors.append(f"slide {row.get('slide_id', '?')}: {path} must declare asset_role=source_cutout")
        if not truthy(record.get("review_only")):
            errors.append(f"slide {row.get('slide_id', '?')}: {path} must declare review_only=true")
        source_hash = str(record.get("source_sha256", "")).strip()
        output_hash = str(record.get("output_sha256", "")).strip()
        if not source_hash or not output_hash:
            errors.append(f"slide {row.get('slide_id', '?')}: {path} lacks source/output hashes")
        if source_hash and not re.fullmatch(r"[0-9a-fA-F]{64}", source_hash):
            errors.append(f"slide {row.get('slide_id', '?')}: {path} has an invalid source_sha256")
        if output_hash and not re.fullmatch(r"[0-9a-fA-F]{64}", output_hash):
            errors.append(f"slide {row.get('slide_id', '?')}: {path} has an invalid output_sha256")
        if not str(record.get("source_hash_basis", "")).strip():
            errors.append(f"slide {row.get('slide_id', '?')}: {path} lacks source_hash_basis")
        figure_id = str(record.get("figure_id", "")).strip()
        if not figure_id:
            errors.append(f"slide {row.get('slide_id', '?')}: source-cutout sidecar {path} lacks figure_id")
        if row_figure_ids and figure_id and figure_id not in row_figure_ids:
            errors.append(f"slide {row.get('slide_id', '?')}: source-cutout figure_id {figure_id!r} is not joined to the manifest")
        rights_basis = str(record.get("rights_basis", "")).strip()
        if not rights_basis:
            errors.append(f"slide {row.get('slide_id', '?')}: source-cutout sidecar {path} lacks rights_basis")
        elif row.get("rights_basis", "").strip() and rights_basis != row.get("rights_basis", "").strip():
            errors.append(f"slide {row.get('slide_id', '?')}: source-cutout rights_basis does not match the manifest")
        fidelity_status = str(record.get("fidelity_status", "")).strip().lower()
        if fidelity_status not in VISUAL_REVIEW_STATUSES:
            errors.append(f"slide {row.get('slide_id', '?')}: source-cutout sidecar {path} has invalid or missing fidelity_status")
        expected_fidelity = row.get("fidelity_status", "").strip().lower()
        if expected_fidelity and fidelity_status and expected_fidelity != fidelity_status:
            errors.append(f"slide {row.get('slide_id', '?')}: source-cutout fidelity_status does not match the manifest")
        if strict and fidelity_status not in {"passed", "not_required"}:
            errors.append(f"slide {row.get('slide_id', '?')}: source-cutout fidelity_status must be passed or not_required in strict mode")
        source_id = str(record.get("source_id", "")).strip()
        if not source_id:
            errors.append(f"slide {row.get('slide_id', '?')}: source-cutout sidecar {path} lacks source_id")
        if row_source_ids and source_id not in row_source_ids:
            errors.append(f"slide {row.get('slide_id', '?')}: source-cutout source_id {source_id!r} is not joined to the manifest")
        record_claim_ids = set(record_values(record, "claim_ids"))
        if not record_claim_ids:
            errors.append(f"slide {row.get('slide_id', '?')}: source-cutout {asset_id!r} lacks claim_ids")
        if row_claim_ids and not row_claim_ids.intersection(record_claim_ids):
            errors.append(f"slide {row.get('slide_id', '?')}: source-cutout {asset_id!r} shares no claim_id with the manifest")
        output_raw = str(record.get("output", "")).strip()
        if not output_raw:
            errors.append(f"slide {row.get('slide_id', '?')}: source-cutout sidecar {path} lacks output")
        else:
            output_path = resolve_audit_path(output_raw, audit_root)
            if output_path is None or not output_path.is_file():
                relative_to_sidecar = path.parent / output_raw
                if relative_to_sidecar.is_file():
                    output_path = relative_to_sidecar
                else:
                    (errors if strict else warnings).append(
                        f"slide {row.get('slide_id', '?')}: source-cutout output does not exist: {output_raw}"
                    )
            if output_path is not None and output_path.is_file():
                if not path_within_audit_root(output_path, audit_root):
                    (errors if strict else warnings).append(
                        f"slide {row.get('slide_id', '?')}: source-cutout output is outside audit-root: {output_path}"
                    )
                else:
                    actual_output_hash = sha256_file(output_path)
                    if actual_output_hash.casefold() != str(record.get("output_sha256", "")).strip().casefold():
                        errors.append(f"slide {row.get('slide_id', '?')}: source-cutout output_sha256 does not match {output_path}")
        source_material_raw = str(
            record.get("source_snapshot_path")
            or record.get("local_image_path")
            or record.get("local_source_path")
            or ""
        ).strip()
        if source_material_raw:
            source_material_path = resolve_audit_path(source_material_raw, audit_root)
            if source_material_path is None or not source_material_path.is_file():
                relative_to_sidecar = path.parent / source_material_raw
                if relative_to_sidecar.is_file():
                    source_material_path = relative_to_sidecar
                else:
                    (errors if strict else warnings).append(
                        f"slide {row.get('slide_id', '?')}: source hash material does not exist: {source_material_raw}"
                    )
            if source_material_path is not None and source_material_path.is_file() and not path_within_audit_root(source_material_path, audit_root):
                (errors if strict else warnings).append(
                    f"slide {row.get('slide_id', '?')}: source hash material is outside audit-root: {source_material_path}"
                )
            elif source_material_path is not None and source_material_path.is_file():
                actual_source_hash = sha256_file(source_material_path)
                if actual_source_hash.casefold() != str(record.get("source_sha256", "")).strip().casefold():
                    errors.append(f"slide {row.get('slide_id', '?')}: source_sha256 does not match {source_material_path}")
    if len(record_ids) != len(set(record_ids)):
        errors.append(f"slide {row.get('slide_id', '?')}: source-cutout sidecar IDs are not unique")
    if set(record_ids) != set(source_cutout_ids):
        errors.append(f"slide {row.get('slide_id', '?')}: source-cutout sidecar IDs do not match source_cutout_asset_ids")
    return errors, warnings, records


def check_scan_records(row, audit_root, strict, expected_route=None):
    """Check a structured no-candidate/access scan record."""
    errors = []
    warnings = []
    records, load_errors, load_warnings = load_json_records(
        row.get("source_scan_record_paths", ""),
        audit_root,
        f"slide {row.get('slide_id', '?')} source-scan record",
        strict,
    )
    errors.extend(load_errors)
    warnings.extend(load_warnings)
    expected_route = expected_route or row.get("source_visual_route", "")
    for _raw, path, record in records:
        expected_visual_id = str(row.get("visual_id", "")).strip()
        record_visual_id = record.get("visual_id", record.get("visualId"))
        if not expected_visual_id:
            errors.append(f"slide {row.get('slide_id', '?')}: source-scan validation requires manifest visual_id")
        if not isinstance(record_visual_id, str) or not record_visual_id.strip():
            errors.append(f"slide {row.get('slide_id', '?')}: {path} lacks visual_id")
        elif expected_visual_id and record_visual_id.strip() != expected_visual_id:
            errors.append(f"slide {row.get('slide_id', '?')}: {path} visual_id does not match the manifest")
        expected_generation_id = str(row.get("generation_id", "")).strip()
        if expected_generation_id:
            record_generation_id = str(record.get("generation_id", record.get("generationId", ""))).strip()
            if not record_generation_id:
                errors.append(f"slide {row.get('slide_id', '?')}: {path} lacks generation_id")
            elif record_generation_id != expected_generation_id:
                errors.append(f"slide {row.get('slide_id', '?')}: {path} generation_id does not match the manifest")
        elif str(record.get("generation_id", record.get("generationId", ""))).strip():
            errors.append(f"slide {row.get('slide_id', '?')}: {path} generation_id is not joined to the manifest")
        record_type = str(record.get("record_type", "")).strip()
        expected_record_type = "source_figure_access" if expected_route == "inaccessible_source" else "source_figure_scan"
        if record_type != expected_record_type:
            errors.append(f"slide {row.get('slide_id', '?')}: {path} must declare record_type={expected_record_type}")
        record_route = str(record.get("source_visual_route", "")).strip()
        if not record_route:
            errors.append(f"slide {row.get('slide_id', '?')}: {path} lacks source_visual_route")
        elif record_route != expected_route:
            errors.append(f"slide {row.get('slide_id', '?')}: {path} route does not match {expected_route}")
        status = str(record.get("status", "")).strip().lower()
        expected_status = "inaccessible" if expected_route == "inaccessible_source" else "no_candidate"
        if status != expected_status:
            errors.append(f"slide {row.get('slide_id', '?')}: {path} must declare status={expected_status}")
        locator = record.get("locator") or record.get("source_locator") or record.get("access_locator")
        if not isinstance(locator, str) or not locator.strip():
            errors.append(f"slide {row.get('slide_id', '?')}: {path} lacks a source/access locator")
        if expected_route == "inaccessible_source":
            access_status = record.get("access_status")
            if not isinstance(access_status, str) or not access_status.strip():
                errors.append(f"slide {row.get('slide_id', '?')}: {path} lacks access_status")
        else:
            reason = record.get("reason") or record.get("summary") or record.get("note")
            if not isinstance(reason, str) or not reason.strip():
                errors.append(f"slide {row.get('slide_id', '?')}: {path} lacks a no-candidate reason")
        observed_at = str(record.get("observed_at") or record.get("created_at") or "").strip()
        if not observed_at:
            errors.append(f"slide {row.get('slide_id', '?')}: {path} lacks observed_at/created_at")
        elif not valid_timestamp(observed_at):
            errors.append(f"slide {row.get('slide_id', '?')}: {path} has an invalid observed_at/created_at")
    return errors, warnings, records


def check_comparison_packet(row, audit_root, strict):
    """Check the structured source/reconstruction comparison packet join."""
    errors = []
    warnings = []
    records, load_errors, load_warnings = load_json_records(
        row.get("visual_comparison_packet_path", ""),
        audit_root,
        f"slide {row.get('slide_id', '?')} visual comparison packet",
        strict,
    )
    errors.extend(load_errors)
    warnings.extend(load_warnings)
    expected_sources = set(pipe_values(row.get("source_cutout_asset_ids", "")))
    expected_derived = str(row.get("derived_asset_id", "")).strip()
    expected_reason = str(row.get("decision_reason_code", "")).strip()
    expected_claims = set(pipe_values(row.get("claim_ids", "")))
    expected_visual_id = str(row.get("visual_id", "")).strip()
    expected_generation_id = str(row.get("generation_id", "")).strip()
    if len(records) != 1:
        errors.append(f"slide {row.get('slide_id', '?')}: visual comparison packet must resolve to exactly one record")
    for _raw, path, packet in records:
        packet_visual_id = str(packet.get("visual_id", packet.get("visualId", ""))).strip()
        if not expected_visual_id:
            errors.append(f"slide {row.get('slide_id', '?')}: comparison packet validation requires manifest visual_id")
        if not packet_visual_id:
            errors.append(f"slide {row.get('slide_id', '?')}: comparison packet {path} lacks visual_id")
        elif expected_visual_id and packet_visual_id != expected_visual_id:
            errors.append(f"slide {row.get('slide_id', '?')}: comparison packet visual_id does not match the manifest")
        if expected_generation_id:
            packet_generation_id = str(packet.get("generation_id", packet.get("generationId", ""))).strip()
            if not packet_generation_id:
                errors.append(f"slide {row.get('slide_id', '?')}: comparison packet {path} lacks generation_id")
            elif packet_generation_id != expected_generation_id:
                errors.append(f"slide {row.get('slide_id', '?')}: comparison packet generation_id does not match the manifest")
        elif str(packet.get("generation_id", packet.get("generationId", ""))).strip():
            errors.append(f"slide {row.get('slide_id', '?')}: comparison packet {path} generation_id is not joined to the manifest")
        packet_sources = set(record_values(packet, "source_cutout_asset_ids"))
        if packet_sources != expected_sources:
            errors.append(f"slide {row.get('slide_id', '?')}: comparison packet source IDs do not match the manifest")
        if str(packet.get("derived_asset_id", "")).strip() != expected_derived:
            errors.append(f"slide {row.get('slide_id', '?')}: comparison packet derived_asset_id does not match the manifest")
        if str(packet.get("decision_reason_code", "")).strip() != expected_reason:
            errors.append(f"slide {row.get('slide_id', '?')}: comparison packet decision_reason_code does not match the manifest")
        packet_status = str(packet.get("visual_review_status", "")).strip().lower()
        if packet_status != str(row.get("visual_review_status", "")).strip().lower():
            errors.append(f"slide {row.get('slide_id', '?')}: comparison packet visual_review_status does not match the manifest")
        packet_claims = set(record_values(packet, "claim_ids"))
        if packet_claims != expected_claims:
            errors.append(f"slide {row.get('slide_id', '?')}: comparison packet claim_ids do not match the manifest")
        for key in ("preserved_relation", "changed_elements", "added_interpretation"):
            if not isinstance(packet.get(key), str) or not packet.get(key, "").strip():
                errors.append(f"slide {row.get('slide_id', '?')}: comparison packet lacks {key}")
        if strict and packet_status != "passed":
            errors.append(f"slide {row.get('slide_id', '?')}: comparison packet must declare visual_review_status=passed in strict mode")
    return errors, warnings


def color_from_element(el: ET.Element):
    for child in list(el):
        if lname(child.tag) == "solidFill":
            for c in child.iter():
                if lname(c.tag) == "srgbClr" and c.attrib.get("val"):
                    return c.attrib["val"].upper()
    return None


def text_and_metrics_from_pptx(path: Path, internal_tokens=INTERNAL_TOKENS):
    slides = []
    color_pages = defaultdict(set)
    sizes = []
    autofit = []
    with zipfile.ZipFile(path) as zf:
        names = sorted(
            [n for n in zf.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)],
            key=natural_key,
        )
        for idx, name in enumerate(names, start=1):
            root = ET.fromstring(zf.read(name))
            text_parts = []
            slide_sizes = []
            slide_colors = set()
            has_autofit = False
            blips = 0
            pictures = 0
            graphic_frames = 0
            connectors = 0
            vector_shapes = 0
            for el in root.iter():
                tag = lname(el.tag)
                if tag == "t" and el.text:
                    text_parts.append(el.text)
                if tag == "blip":
                    blips += 1
                if tag == "pic":
                    pictures += 1
                if tag == "graphicFrame":
                    graphic_frames += 1
                if tag == "cxnSp":
                    connectors += 1
                if tag == "sp":
                    shape_props = next(
                        (child for child in el if lname(child.tag) == "spPr"), None
                    )
                    has_fill_or_line = bool(
                        shape_props is not None
                        and any(
                            lname(child.tag) in {"solidFill", "gradFill", "pattFill", "blipFill", "ln"}
                            for child in shape_props.iter()
                        )
                    )
                    # A styled shape can be a labelled rectangle or a vector
                    # diagram element. A plain text box is not counted as a
                    # visual object; semantic adequacy still needs human review.
                    if has_fill_or_line:
                        vector_shapes += 1
                if tag in {"normAutofit", "spAutoFit", "prstTxWarp"}:
                    has_autofit = True
                if tag in {"rPr", "defRPr", "endParaRPr"}:
                    if el.attrib.get("sz"):
                        try:
                            pt = int(el.attrib["sz"]) / 100.0
                            slide_sizes.append(pt)
                            sizes.append(pt)
                        except ValueError:
                            pass
                    c = color_from_element(el)
                    if c:
                        slide_colors.add(c)
                        color_pages[c].add(idx)
            if has_autofit:
                autofit.append(idx)
            observed_visuals = max(blips, pictures) + graphic_frames + connectors + vector_shapes
            text = " ".join(text_parts)
            slides.append({
                "slide": idx,
                "text": text,
                "font_sizes_pt": slide_sizes,
                "text_colors": sorted(slide_colors),
                "embedded_images": max(blips, pictures),
                "graphic_frames": graphic_frames,
                "connectors": connectors,
                "vector_shapes": vector_shapes,
                "non_text_shapes": vector_shapes,
                "observed_visuals": observed_visuals,
                "has_observed_visual": observed_visuals > 0,
                "citation": citation_match(text),
                "numeric_tokens": numeric_matches(text),
                "body_candidate_font_sizes_pt": [],
                "body_candidate_text": 0,
                "autofit": has_autofit,
            })
    try:
        from pptx import Presentation

        presentation = Presentation(str(path))
        for index, slide in enumerate(presentation.slides):
            candidates = []
            body_candidate_text = 0
            for shape in slide.shapes:
                if not getattr(shape, "has_text_frame", False):
                    continue
                shape_text = shape.text.strip()
                if len(shape_text) < 12 or (citation_match(shape_text) and is_provenance_only(shape_text)):
                    continue
                body_candidate_text += 1
                sizes_for_shape = []
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.font.size is not None:
                            sizes_for_shape.append(run.font.size.pt)
                if sizes_for_shape:
                    candidates.extend(sizes_for_shape)
            if index < len(slides):
                slides[index]["body_candidate_font_sizes_pt"] = candidates
                slides[index]["body_candidate_text"] = body_candidate_text
    except Exception:
        # Explicit XML sizes remain useful; inherited or unavailable shape
        # information is reported as a review warning rather than a hard fail.
        pass
    internal = []
    for slide in slides:
        for token in internal_tokens:
            if token.casefold() in slide["text"].casefold():
                internal.append({"slide": slide["slide"], "token": token})
    return {
        "slides": slides,
        "slide_count": len(slides),
        "text_colors": sorted(color_pages),
        "accent_pages": sorted({p for c, pages in color_pages.items() if c in DEFAULT_ACCENTS for p in pages}),
        "font_sizes_pt": sizes,
        "autofit_slides": autofit,
        "internal_tokens": internal,
    }


def pptx_geometry(path: Path):
    """Return conservative text-text collision and bounds diagnostics."""
    try:
        from pptx import Presentation
    except Exception as exc:  # pragma: no cover - depends on runtime
        return {"available": False, "warning": f"python-pptx unavailable: {exc}"}
    prs = Presentation(str(path))
    collisions = []
    out_of_bounds = []
    sw, sh = prs.slide_width, prs.slide_height
    for i, slide in enumerate(prs.slides, start=1):
        text_shapes = []
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            text = shape.text.strip()
            if not text:
                continue
            x1, y1 = shape.left, shape.top
            x2, y2 = x1 + shape.width, y1 + shape.height
            if x1 < 0 or y1 < 0 or x2 > sw or y2 > sh:
                out_of_bounds.append({"slide": i, "text": text[:80]})
            text_shapes.append((shape, x1, y1, x2, y2, text))
        for a in range(len(text_shapes)):
            _, ax1, ay1, ax2, ay2, at = text_shapes[a]
            for b in range(a + 1, len(text_shapes)):
                _, bx1, by1, bx2, by2, bt = text_shapes[b]
                iw = max(0, min(ax2, bx2) - max(ax1, bx1))
                ih = max(0, min(ay2, by2) - max(ay1, by1))
                inter = iw * ih
                if inter and inter / min((ax2 - ax1) * (ay2 - ay1), (bx2 - bx1) * (by2 - by1)) > 0.15:
                    collisions.append({"slide": i, "a": at[:60], "b": bt[:60]})
    return {"available": True, "collisions": collisions, "out_of_bounds": out_of_bounds}


def read_manifest(
    path: Path | None,
    strict: bool = False,
    audit_root: Path | None = None,
    plot_required: bool = False,
):
    if not path:
        return None, ["visual_manifest.tsv was not supplied"], []
    rows = []
    errors = []
    warnings = []
    required = {
        "slide_id", "narrative_job", "evidence_class",
        "citation_visible", "numeric", "denominator_unit",
        "aggregation_level", "caveat_visible", "min_body_pt",
        "protected_terms", "internal_vocab_free",
    }
    if plot_required:
        required.update({"plot_version", "plot_row_id", "section_id", "evidence_block_id"})
    visual_contract_columns = {
        "visual_id",
        "kind",
        "claim_bearing",
        "contract_path",
        "source_ids",
        "claim_ids",
        "rights_basis",
        "source_cutout_asset_ids",
        "source_cutout_manifest_paths",
        "source_scan_record_paths",
        "derived_asset_id",
        "derived_asset_path",
        "decision_reason_code",
        "attribution_note",
        "visual_review_status",
        "semantic_status",
        "fidelity_status",
        "reader_role",
        "boundary",
        "visual_comparison_packet_path",
        "source_cutout_created_at",
        "reconstruction_created_at",
    }
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream, delimiter="\t")
            missing = required - set(reader.fieldnames or [])
            if missing:
                errors.append(f"manifest missing columns: {sorted(missing)}")
            for row in reader:
                if None in row:
                    errors.append("manifest contains extra TSV fields")
                rows.append(row)
            if any(not has_nested_visuals(row) for row in rows):
                legacy_required = {"primary_visual_type", "source_pointer", "source_visual_route", "figure_id", "status"}
                missing_legacy = legacy_required - set(reader.fieldnames or [])
                if missing_legacy:
                    errors.append(f"manifest missing columns: {sorted(missing_legacy)}")
                missing_visual = visual_contract_columns - set(reader.fieldnames or [])
                if missing_visual:
                    message = f"manifest missing claim-bearing visual contract columns: {sorted(missing_visual)}"
                    (errors if strict else warnings).append(message)
    except (OSError, UnicodeError, csv.Error) as exc:
        errors.append(f"manifest could not be read: {exc}")
        return rows, errors, warnings
    visual_rows, visual_errors = deck_visual_rows(rows)
    errors.extend(visual_errors)
    for row in visual_rows:
        sid = row.get("slide_id", "?")
        if not SLIDE_ID_PATTERN.fullmatch(sid or ""):
            errors.append(f"slide {sid}: slide_id must be numeric or use S1/slide-1 form")
        visual = row.get("primary_visual_type", "")
        cls = row.get("evidence_class", "")
        if visual not in ALLOWED_VISUALS:
            errors.append(f"slide {sid}: invalid primary_visual_type={visual!r}")
        if visual == "none" and cls not in NON_EVIDENCE_CLASSES:
            errors.append(f"slide {sid}: evidence class {cls!r} cannot use none")
        source_visual_route = row.get("source_visual_route", "")
        if source_visual_route not in SOURCE_VISUAL_ROUTES:
            errors.append(f"slide {sid}: invalid source_visual_route={source_visual_route!r}")
        source_cutout_ids = pipe_values(row.get("source_cutout_asset_ids", ""))
        source_cutout_manifests = pipe_values(row.get("source_cutout_manifest_paths", ""))
        source_scan_records = pipe_values(row.get("source_scan_record_paths", ""))
        reason = row.get("decision_reason_code", "").strip()
        review_status = row.get("visual_review_status", "").strip().lower()
        if review_status and review_status not in VISUAL_REVIEW_STATUSES:
            errors.append(f"slide {sid}: invalid visual_review_status={review_status!r}")
        if source_visual_route == "source_figure":
            if visual != "source_figure":
                errors.append(f"slide {sid}: source_visual_route=source_figure requires primary_visual_type=source_figure")
            if source_scan_records:
                errors.append(f"slide {sid}: source_figure route must not use source_scan_record_paths")
            if not row.get("figure_id", "").strip():
                errors.append(f"slide {sid}: source_figure route lacks figure_id")
            if not source_cutout_ids:
                errors.append(f"slide {sid}: source_figure route lacks source_cutout_asset_ids")
            if not source_cutout_manifests:
                errors.append(f"slide {sid}: source_figure route lacks source_cutout_manifest_paths")
            if not pipe_values(row.get("source_ids", "")):
                errors.append(f"slide {sid}: source_figure route lacks source_ids")
            if not pipe_values(row.get("claim_ids", "")):
                errors.append(f"slide {sid}: source_figure route lacks claim_ids")
            if not valid_timestamp(row.get("source_cutout_created_at", "")):
                errors.append(f"slide {sid}: source_figure route lacks a valid source_cutout_created_at")
            if strict and review_status not in {"passed", "not_required"}:
                errors.append(f"slide {sid}: source_figure route requires visual_review_status=not_required or passed in strict mode")
            rights_basis = row.get("rights_basis", "").strip().lower()
            if strict and rights_basis in {"", "not provided", "unclear"}:
                errors.append(f"slide {sid}: source_figure route lacks an established rights_basis")
            source_errors, source_warnings, _records = check_source_cutout_records(
                row, source_cutout_ids, source_cutout_manifests, audit_root, strict
            )
            errors.extend(source_errors)
            warnings.extend(source_warnings)
        elif source_visual_route in RECONSTRUCTION_ROUTES:
            if visual == "none":
                errors.append(f"slide {sid}: reconstruction route cannot use primary_visual_type=none")
            if not row.get("derived_asset_id", "").strip():
                errors.append(f"slide {sid}: reconstruction route lacks derived_asset_id")
            if not row.get("derived_asset_path", "").strip():
                errors.append(f"slide {sid}: reconstruction route lacks derived_asset_path")
            if reason not in NON_USE_REASON_CODES:
                errors.append(f"slide {sid}: reconstruction route has invalid decision_reason_code={reason!r}")
            if source_visual_route == "new_explanatory_visual" and reason == "no_source_figure":
                errors.append(
                    f"slide {sid}: new_explanatory_visual cannot use no_source_figure; use no_useful_candidate route"
                )
            if not row.get("attribution_note", "").strip():
                errors.append(f"slide {sid}: reconstruction route lacks reader-plane attribution_note")
            if not row.get("visual_comparison_packet_path", "").strip():
                errors.append(f"slide {sid}: reconstruction route lacks visual_comparison_packet_path")
            if reason != "no_source_figure" and not valid_timestamp(row.get("source_cutout_created_at", "")):
                errors.append(f"slide {sid}: reconstruction route lacks a valid source_cutout_created_at")
            if not valid_timestamp(row.get("reconstruction_created_at", "")):
                errors.append(f"slide {sid}: reconstruction route lacks a valid reconstruction_created_at")
            no_source_faithful = source_visual_route == "faithful_reconstruction" and reason == "no_source_figure"
            if not no_source_faithful and not pipe_values(row.get("source_ids", "")):
                errors.append(f"slide {sid}: reconstruction route lacks source_ids")
            if not pipe_values(row.get("claim_ids", "")):
                errors.append(f"slide {sid}: reconstruction route lacks claim_ids")
            source_time = timestamp_value(row.get("source_cutout_created_at", ""))
            reconstruction_time = timestamp_value(row.get("reconstruction_created_at", ""))
            if source_time and reconstruction_time and reconstruction_time < source_time:
                errors.append(f"slide {sid}: reconstruction_created_at precedes source_cutout_created_at")
            derived_missing, derived_uncheckable = audit_paths_exist(
                row.get("derived_asset_path", ""), audit_root
            )
            derived_outside = audit_paths_outside_root(row.get("derived_asset_path", ""), audit_root)
            if derived_missing:
                (errors if strict else warnings).append(f"slide {sid}: missing derived asset paths: {derived_missing}")
            if derived_uncheckable:
                (errors if strict else warnings).append(f"slide {sid}: derived asset URLs were not checked locally: {derived_uncheckable}")
            if derived_outside:
                (errors if strict else warnings).append(f"slide {sid}: derived asset paths are outside audit-root: {derived_outside}")
            if reason == "no_source_figure":
                if source_cutout_ids or source_cutout_manifests:
                    errors.append(f"slide {sid}: no_source_figure must use source_scan_record_paths, not source-cutout paths")
                if not source_scan_records:
                    errors.append(f"slide {sid}: no_source_figure requires source_scan_record_paths")
                scan_errors, scan_warnings, scan_records = check_scan_records(
                    row, audit_root, strict, expected_route="no_useful_candidate"
                )
                errors.extend(scan_errors)
                warnings.extend(scan_warnings)
                if reconstruction_time:
                    for _raw, _path, scan_record in scan_records:
                        scan_time = timestamp_value(str(scan_record.get("observed_at") or scan_record.get("created_at") or ""))
                        if scan_time and reconstruction_time < scan_time:
                            errors.append(f"slide {sid}: reconstruction_created_at precedes source-scan observation")
            else:
                if not source_cutout_ids:
                    errors.append(f"slide {sid}: reconstruction route lacks source_cutout_asset_ids")
                if not source_cutout_manifests:
                    errors.append(f"slide {sid}: reconstruction route lacks source_cutout_manifest_paths")
                if not pipe_values(row.get("source_ids", "")):
                    errors.append(f"slide {sid}: reconstruction route lacks source_ids")
                if not pipe_values(row.get("claim_ids", "")):
                    errors.append(f"slide {sid}: reconstruction route lacks claim_ids")
                source_errors, source_warnings, _records = check_source_cutout_records(
                    row, source_cutout_ids, source_cutout_manifests, audit_root, strict
                )
                errors.extend(source_errors)
                warnings.extend(source_warnings)
            if reason == "cross_source_comparison" and len(source_cutout_ids) < 2:
                errors.append(f"slide {sid}: cross_source_comparison requires at least two source_cutout_asset_ids")
            if reason == "cross_source_comparison" and len(set(pipe_values(row.get("source_ids", "")))) < 2:
                errors.append(f"slide {sid}: cross_source_comparison requires at least two source_ids")
            if strict and review_status != "passed":
                errors.append(f"slide {sid}: reconstruction route requires visual_review_status=passed in strict mode")
            rights_basis = row.get("rights_basis", "").strip().lower()
            if strict and source_cutout_ids and rights_basis in {"", "not provided", "unclear"}:
                errors.append(f"slide {sid}: reconstruction route lacks an established rights_basis for its source cutout")
            packet_errors, packet_warnings = check_comparison_packet(row, audit_root, strict)
            errors.extend(packet_errors)
            warnings.extend(packet_warnings)
        elif source_visual_route == "no_useful_candidate":
            if reason != "no_source_figure":
                errors.append(f"slide {sid}: no_useful_candidate requires decision_reason_code=no_source_figure")
            if source_cutout_ids or source_cutout_manifests:
                errors.append(f"slide {sid}: no_useful_candidate must not use source-cutout paths")
            if not source_scan_records:
                errors.append(f"slide {sid}: no_useful_candidate lacks source_scan_record_paths")
            scan_errors, scan_warnings, _scan_records = check_scan_records(row, audit_root, strict)
            errors.extend(scan_errors)
            warnings.extend(scan_warnings)
        elif source_visual_route == "inaccessible_source":
            if not source_scan_records:
                errors.append(f"slide {sid}: inaccessible_source lacks source_scan_record_paths")
            scan_errors, scan_warnings, _scan_records = check_scan_records(row, audit_root, strict)
            errors.extend(scan_errors)
            warnings.extend(scan_warnings)
            if strict and cls in EVIDENCE_CLASSES:
                errors.append(f"slide {sid}: inaccessible_source is unresolved and cannot pass strict release")
        if cls in EVIDENCE_CLASSES:
            if source_visual_route == "not_applicable":
                errors.append(f"slide {sid}: evidence page must record a G4-V source-visual route")
            if not row.get("source_pointer", "").strip():
                errors.append(f"slide {sid}: evidence page lacks a source_pointer")
    for row in rows:
        sid = row.get("slide_id", "?")
        cls = row.get("evidence_class", "")
        if cls in EVIDENCE_CLASSES:
            if row.get("citation_visible", "").lower() not in {"yes", "true", "1"}:
                errors.append(f"slide {sid}: evidence page must plan a visible citation")
            if row.get("internal_vocab_free", "").lower() not in {"yes", "true", "1"}:
                errors.append(f"slide {sid}: internal vocabulary not cleared in the plan")
            try:
                float(row.get("min_body_pt", "0"))
            except ValueError:
                errors.append(f"slide {sid}: invalid min_body_pt")
            numeric = row.get("numeric", "").lower()
            if numeric in {"yes", "true", "1"} and not row.get("denominator_unit", "").strip():
                errors.append(f"slide {sid}: numeric page lacks denominator/unit in the plan")
    page_numbers = [slide_number(row.get("slide_id", "")) for row in rows]
    valid_page_numbers = [number for number in page_numbers if number is not None]
    if len(valid_page_numbers) != len(set(valid_page_numbers)):
        errors.append("manifest contains duplicate slide/page mappings")
    return rows, errors, warnings


def check_visual_contract_links(rows, audit_root, strict):
    """Check the G6 artifact manifest's joins to per-visual contracts.

    The structured evidence checker owns deep semantic/route validation. This
    smaller observer check prevents a deck manifest from silently pointing at
    a different visual, kind, route, or review state than the contract that was
    supposedly built.
    """
    rows, errors = deck_visual_rows(rows)
    warnings = []
    if not rows:
        return errors, warnings
    contract_columns = {
        "visual_id", "kind", "claim_bearing", "contract_path",
        "semantic_status", "fidelity_status", "visual_review_status",
        "reader_role", "boundary", "status",
    }
    seen_ids = set()
    for row in rows:
        sid = row.get("slide_id", "?")
        claim_ids = pipe_values(row.get("claim_ids", ""))
        contract_raw = row.get("contract_path", "").strip()
        present = set(row)
        contract_mode = bool(present & {"visual_id", "kind", "claim_bearing", "contract_path"})
        if claim_ids or truthy(row.get("claim_bearing", "")):
            missing = contract_columns - present
            if missing:
                message = f"slide {sid} visual {row.get('visual_id', '?')}: claim-bearing row requires contract columns: {sorted(missing)}"
                (errors if strict else warnings).append(message)
        if not claim_ids and not contract_raw:
            if contract_mode:
                claim_bearing_raw = row.get("claim_bearing", "")
                claim_bearing_text = claim_bearing_raw.strip().lower()
                if claim_bearing_text not in {"true", "false", "yes", "no", "1", "0"}:
                    errors.append(f"slide {sid}: claim_bearing must be explicit even when claim_ids and contract_path are empty")
                elif truthy(claim_bearing_raw):
                    errors.append(f"slide {sid}: claim-bearing row lacks claim_ids and contract_path")
            continue
        visual_id = row.get("visual_id", "").strip()
        if not visual_id:
            errors.append(f"slide {sid}: claim-bearing contract join lacks visual_id")
            continue
        if visual_id in seen_ids:
            errors.append(f"manifest contains duplicate visual_id={visual_id}")
        seen_ids.add(visual_id)
        kind = row.get("kind", "").strip()
        claim_bearing_raw = row.get("claim_bearing", "")
        claim_bearing_text = claim_bearing_raw.strip().lower()
        if claim_bearing_text not in {"true", "false", "yes", "no", "1", "0"}:
            errors.append(f"slide {sid}: claim_bearing must be an explicit boolean")
            claim_bearing = bool(claim_ids)
        else:
            claim_bearing = truthy(claim_bearing_raw)
        if claim_ids and not claim_bearing:
            errors.append(f"slide {sid}: claim_ids are present but claim_bearing is false")
        if claim_bearing and kind not in EVIDENCE_CONTRACT_KINDS:
            errors.append(f"slide {sid}: unsupported claim-bearing contract kind={kind!r}")
        for status_name in ("semantic_status", "fidelity_status", "visual_review_status"):
            value = row.get(status_name, "").strip().lower()
            if value and value not in VISUAL_REVIEW_STATUSES:
                errors.append(f"slide {sid}: invalid {status_name}={value!r}")
        semantic_status = row.get("semantic_status", "").strip().lower()
        if claim_bearing and not semantic_status:
            errors.append(f"slide {sid}: claim-bearing row lacks semantic_status")
        if strict and claim_bearing and semantic_status != "passed":
            errors.append(f"slide {sid}: claim-bearing row requires semantic_status=passed in strict mode")
        fidelity_status = row.get("fidelity_status", "").strip().lower()
        route = row.get("source_visual_route", "").strip()
        reason = row.get("decision_reason_code", "").strip()
        if route == "source_figure" and kind != "source_figure":
            errors.append(f"slide {sid}: source_visual_route=source_figure requires kind=source_figure")
        if strict and claim_bearing and route == "source_figure" and fidelity_status not in {"passed", "not_required"}:
            errors.append(f"slide {sid}: source_figure route requires fidelity_status=passed or not_required in strict mode")
        if strict and claim_bearing and route == "faithful_reconstruction":
            allowed_fidelity = {"passed", "not_required"} if reason == "no_source_figure" else {"passed"}
            if fidelity_status not in allowed_fidelity:
                errors.append(f"slide {sid}: source-backed faithful_reconstruction requires fidelity_status=passed in strict mode")
        if not contract_raw:
            errors.append(f"slide {sid}: claim-bearing row lacks contract_path")
            continue
        contract_path = resolve_audit_path(contract_raw, audit_root)
        if contract_path is None or not path_within_audit_root(contract_path, audit_root):
            errors.append(f"slide {sid}: contract_path is not local to audit-root: {contract_raw}")
            continue
        if not contract_path.is_file():
            errors.append(f"slide {sid}: contract_path does not exist: {contract_path}")
            continue
        try:
            contract = json.loads(contract_path.read_text(encoding="utf-8"))
        except (OSError, UnicodeError, json.JSONDecodeError) as exc:
            errors.append(f"slide {sid}: contract_path is not valid JSON: {exc}")
            continue
        if not isinstance(contract, dict):
            errors.append(f"slide {sid}: contract_path must contain one JSON object")
            continue
        review = contract.get("review")
        if not isinstance(review, dict):
            errors.append(f"slide {sid}: contract requires nested review status object")
            review = {}
        for status_name in ("deterministic_status", "semantic_status", "fidelity_status"):
            camel_name = {
                "deterministic_status": "deterministicStatus",
                "semantic_status": "semanticStatus",
                "fidelity_status": "fidelityStatus",
            }[status_name]
            nested_status = str(review.get(status_name, review.get(camel_name, ""))).strip().lower()
            if nested_status not in VISUAL_REVIEW_STATUSES:
                errors.append(f"slide {sid}: contract.review.{status_name} has invalid or missing status")
            if strict and status_name in {"deterministic_status", "semantic_status"} and nested_status != "passed":
                errors.append(f"slide {sid}: contract.review.{status_name} must be passed in strict mode")
            direct_status = str(contract.get(status_name, "")).strip().lower()
            if direct_status and nested_status and direct_status != nested_status:
                errors.append(f"slide {sid}: contract.{status_name} does not match contract.review.{status_name}")
        checks = {
            "visual_id": visual_id,
            "slide_id": sid,
            "kind": kind,
            "claim_bearing": claim_bearing,
            "source_visual_route": row.get("source_visual_route", "").strip(),
            "reader_role": row.get("reader_role", "").strip(),
            "boundary": row.get("boundary", "").strip(),
            "status": row.get("status", "").strip(),
        }
        for key, expected in checks.items():
            actual = contract.get(key)
            if key == "claim_bearing":
                if not isinstance(actual, bool) or actual != expected:
                    errors.append(f"slide {sid}: contract.{key} does not match manifest")
            elif str(actual or "").strip() != expected:
                errors.append(f"slide {sid}: contract.{key} does not match manifest")
        for status_name in ("semantic_status", "fidelity_status", "visual_review_status"):
            expected = row.get(status_name, "").strip().lower()
            actual = str(contract.get(status_name, review.get(status_name, ""))).strip().lower()
            if expected and actual != expected:
                errors.append(f"slide {sid}: contract.{status_name} does not match manifest")
    return errors, warnings


def protected_terms_from_manifest(rows):
    terms = set()
    for row in rows or []:
        for term in re.split(r"[|,]", row.get("protected_terms", "")):
            term = term.strip()
            if term:
                terms.add(term)
    return sorted(terms, key=len, reverse=True)


def pdf_metrics(path: Path, terms):
    try:
        import fitz
    except Exception as exc:  # pragma: no cover
        return {"available": False, "warning": f"PyMuPDF unavailable: {exc}"}
    doc = fitz.open(str(path))
    broken_glyph_pages = []
    line_breaks = []
    protected_splits = []
    pages = []
    for pno, page in enumerate(doc, start=1):
        data = page.get_text("dict")
        lines = []
        page_font_sizes = []
        body_candidate_font_sizes = []
        body_candidate_text = 0
        for block in data.get("blocks", []):
            for line in block.get("lines", []):
                spans = line.get("spans", [])
                text = "".join(span.get("text", "") for span in spans)
                if text:
                    lines.append(text)
                sizes = []
                for span in spans:
                    try:
                        size = float(span.get("size"))
                    except (TypeError, ValueError):
                        continue
                    page_font_sizes.append(size)
                    sizes.append(size)
                if len(text) >= 12 and not (citation_match(text) and is_provenance_only(text)):
                    body_candidate_text += 1
                    body_candidate_font_sizes.extend(sizes)
        text = "\n".join(lines)
        if "\ufffd" in text or "�" in text:
            broken_glyph_pages.append(pno)
        for line in lines:
            if line and (line[0] in FORBIDDEN_LINE_START or line[-1] in FORBIDDEN_LINE_END):
                line_breaks.append({"page": pno, "line": line[:100]})
        for left, right in zip(lines, lines[1:]):
            for term in terms:
                for cut in range(1, len(term)):
                    if left.endswith(term[:cut]) and right.startswith(term[cut:]):
                        protected_splits.append({"page": pno, "term": term})
        visual_metrics = pdf_visual_metrics(page)
        pages.append({
            "page": pno,
            "text": text,
            "citation": citation_match(text),
            "numeric_tokens": numeric_matches(text),
            "observed_visuals": visual_metrics["observed_visuals"],
            "has_observed_visual": visual_metrics["has_observed_visual"],
            "images": visual_metrics["images"],
            "drawings": visual_metrics["drawings"],
            "observed_vector_drawings": visual_metrics["observed_vector_drawings"],
            "ignored_background_drawings": visual_metrics["ignored_background_drawings"],
            "ignored_tiny_drawings": visual_metrics["ignored_tiny_drawings"],
            "font_sizes_pt": page_font_sizes,
            "body_candidate_font_sizes_pt": body_candidate_font_sizes,
            "body_candidate_text": body_candidate_text,
        })
    return {
        "available": True,
        "page_count": len(doc),
        "pages": pages,
        "broken_glyph_pages": broken_glyph_pages,
        "forbidden_line_breaks": line_breaks,
        "protected_term_splits": protected_splits,
    }


def validate_observed_artifact(
    rows, pages, min_visuals=None, min_body_pt=None, strict=False, artifact_label="artifact"
):
    failures = []
    warnings = []
    evidence_rows = [row for row in rows if row.get("evidence_class") in EVIDENCE_CLASSES]
    observed_visual_rows = 0
    explicit_font_sizes = []
    body_candidate_sizes = []
    numeric_pages = []
    for position, row in enumerate(rows, start=1):
        number = slide_number(row.get("slide_id", ""))
        if number is None:
            failures.append(f"{artifact_label}: invalid slide_id {row.get('slide_id', '?')}")
            continue
        page = pages.get(number)
        if page is None:
            failures.append(f"{artifact_label}: manifest row {row.get('slide_id', '?')} maps to missing page {number}")
            continue
        cls = row.get("evidence_class", "")
        visuals, visual_errors = deck_visual_rows([row])
        failures.extend(visual_errors)
        visual_types = [visual.get("primary_visual_type", "") for visual in visuals]
        row_floor = min_body_pt
        if row_floor is None:
            try:
                row_floor = float(row.get("min_body_pt", ""))
            except ValueError:
                row_floor = None
        numeric_tokens = page.get("numeric_tokens", [])
        numeric_declared = row.get("numeric", "").lower() in {"yes", "true", "1"}
        if numeric_tokens:
            numeric_pages.append(number)
            if not numeric_declared:
                failures.append(
                    f"{artifact_label} slide {number}: observed numeric tokens {numeric_tokens[:5]} but manifest numeric is not yes"
                )
        elif numeric_declared:
            warnings.append(f"{artifact_label} slide {number}: numeric page has no text numeric token; inspect visual/table content")
        if cls in EVIDENCE_CLASSES:
            if row.get("citation_visible", "").lower() in {"yes", "true", "1"} and not page.get("citation"):
                failures.append(f"{artifact_label} slide {number}: manifest promises visible citation but no reader-facing citation was observed")
            for visual in visuals:
                if visual.get("source_visual_route", "") in RECONSTRUCTION_ROUTES:
                    attribution_note = visual.get("attribution_note", "").strip()
                    if not attribution_observed(
                        page.get("text", ""), attribution_note, allow_generic=not has_nested_visuals(row)
                    ):
                        failures.append(
                            f"{artifact_label} slide {number} visual {visual.get('visual_id', '?')}: reconstruction route lacks an observed reader-plane attribution note"
                        )
            if any(kind != "none" for kind in visual_types):
                if page.get("has_observed_visual"):
                    observed_visual_rows += 1
                else:
                    failures.append(f"{artifact_label} slide {number}: manifest declares {visual_types} but no image/chart/table/vector object was observed")
        if "font_sizes_pt" in page:
            explicit_font_sizes.extend(page["font_sizes_pt"])
            body_candidate_sizes.extend(page.get("body_candidate_font_sizes_pt", []))
            body_text_observed = page.get("body_candidate_text", 0)
            body_sizes = page.get("body_candidate_font_sizes_pt", [])
            if body_text_observed and not body_sizes and row_floor is not None:
                message = f"{artifact_label} slide {number}: body-candidate text exists but its font size was not observable"
                if strict:
                    failures.append(message)
                else:
                    warnings.append(message)
            if row_floor is not None and body_sizes and min(body_sizes) < row_floor:
                failures.append(f"{artifact_label} slide {number}: observed body-candidate text is below task floor {row_floor:g}pt")
    if evidence_rows and observed_visual_rows == 0:
        failures.append(f"{artifact_label}: no observed evidence-bearing visual objects matched the manifest")
    if min_visuals is not None and observed_visual_rows < min_visuals:
        failures.append(f"{artifact_label}: observed evidence visual-object rows {observed_visual_rows} < task/profile floor {min_visuals}")
    if explicit_font_sizes:
        warnings.append("font size measurement covers explicit artifact properties; inherited theme/master sizes need review")
    if min_body_pt is None and not explicit_font_sizes and any(page.get("text") for page in pages.values()):
        warnings.append("PPTX contains text but no explicit run font sizes were observed; inspect inherited typography")
    return {
        "observed_evidence_visual_object_rows": observed_visual_rows,
        "semantic_meaning_review_required": observed_visual_rows > 0,
        "numeric_pages": numeric_pages,
        "evidence_rows": len(evidence_rows),
        "explicit_font_min_pt": min(explicit_font_sizes) if explicit_font_sizes else None,
        "explicit_font_max_pt": max(explicit_font_sizes) if explicit_font_sizes else None,
        "body_candidate_font_min_pt": min(body_candidate_sizes) if body_candidate_sizes else None,
        "body_candidate_font_max_pt": max(body_candidate_sizes) if body_candidate_sizes else None,
    }, failures, warnings


def _main_impl():
    parser = argparse.ArgumentParser()
    parser.add_argument("--pptx", type=Path)
    parser.add_argument("--pdf", type=Path)
    parser.add_argument("--manifest", type=Path)
    parser.add_argument("--plot", "--storyline-plot", dest="plot", type=Path)
    parser.add_argument("--review", "--review-record", dest="review", type=Path)
    parser.add_argument("--min-observed-visuals", type=int)
    parser.add_argument("--min-body-pt", type=float)
    parser.add_argument(
        "--audit-root",
        type=Path,
        help="audit bundle root; required before referenced source-cutout, scan, derived, or comparison-packet paths are read",
    )
    parser.add_argument(
        "--internal-token",
        action="append",
        default=[],
        help="optional task-specific reader-plane token to reject; repeat as needed",
    )
    parser.add_argument("--strict", action="store_true")
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()

    report = {"failures": [], "warnings": [], "metrics": {}}
    audit_root = args.audit_root
    current_plot_version = None
    plot_required = bool(args.strict and args.manifest)
    if plot_required and not args.plot:
        report["failures"].append("strict presentation check requires --plot/--storyline-plot")
    review_required = bool(args.strict and args.manifest)
    if review_required and not args.review:
        report["failures"].append("strict presentation check requires --review/--review-record")
    rows, manifest_errors, manifest_warnings = read_manifest(
        args.manifest,
        strict=args.strict,
        audit_root=audit_root,
        plot_required=plot_required,
    )
    report["metrics"]["manifest_rows"] = len(rows or [])
    visual_rows, _visual_errors = deck_visual_rows(rows or [])
    report["metrics"]["manifest_visuals"] = len(visual_rows)
    report["failures"].extend(manifest_errors)
    report["warnings"].extend(manifest_warnings)
    if args.plot:
        if not args.plot.exists():
            report["failures"].append(f"missing storyline plot: {args.plot}")
        else:
            plot_errors, plot_warnings, plot_metrics = check_plot_file(
                args.plot,
                rows or [],
                strict=args.strict,
            )
            report["metrics"]["storyline_plot"] = plot_metrics
            current_plot_version = plot_metrics.get("plot_version")
            report["failures"].extend(plot_errors)
            report["warnings"].extend(plot_warnings)
    if args.review:
        if not args.review.exists():
            report["failures"].append(f"missing review record: {args.review}")
        else:
            review_plot_rows = []
            if args.plot and args.plot.exists():
                plot_payload, plot_load_errors = _load_payload(args.plot)
                if not plot_load_errors:
                    review_plot_rows = plot_rows_from_payload(plot_payload)
            review_errors, review_warnings, review_metrics = check_review_file(
                args.review,
                review_plot_rows,
                rows or [],
                strict=args.strict,
                expected_plot_version=current_plot_version,
            )
            report["metrics"]["review_record"] = review_metrics
            report["failures"].extend(review_errors)
            report["warnings"].extend(review_warnings)
    contract_errors, contract_warnings = check_visual_contract_links(rows or [], audit_root, args.strict)
    report["failures"].extend(contract_errors)
    report["warnings"].extend(contract_warnings)
    if not args.pptx and not args.pdf:
        report["failures"].append("supply a PPTX or PDF so the manifest can be checked against an artifact")

    artifact_pages = []
    ppt = None
    if args.pptx:
        if not args.pptx.exists():
            report["failures"].append(f"missing PPTX: {args.pptx}")
        else:
            ppt = text_and_metrics_from_pptx(args.pptx, tuple(INTERNAL_TOKENS) + tuple(args.internal_token))
            report["metrics"]["pptx"] = ppt
            artifact_pages.append(("PPTX", {slide["slide"]: slide for slide in ppt["slides"]}))
            geometry = pptx_geometry(args.pptx)
            report["metrics"]["geometry"] = geometry
            if geometry.get("available") is False:
                report["warnings"].append(geometry["warning"])
            elif geometry.get("collisions"):
                report["failures"].append(f"text collisions: {len(geometry['collisions'])}")
            if geometry.get("out_of_bounds"):
                report["failures"].append(f"out-of-bounds text boxes: {len(geometry['out_of_bounds'])}")
            if ppt["autofit_slides"]:
                report["failures"].append(f"auto-fit detected on slides {ppt['autofit_slides']}")
            if ppt["internal_tokens"]:
                report["failures"].append(f"internal reader-plane tokens: {ppt['internal_tokens'][:12]}")
            if ppt["slide_count"] and rows is not None and len(rows) != ppt["slide_count"]:
                report["failures"].append(f"manifest rows {len(rows)} != PPTX slides {ppt['slide_count']}")
            if ppt["text_colors"] and len(ppt["text_colors"]) > 3:
                report["failures"].append(f"authoring text colors exceed 3: {ppt['text_colors']}")
            if ppt["slide_count"]:
                accent_share = len(ppt["accent_pages"]) / ppt["slide_count"]
                report["metrics"]["accent_page_share"] = accent_share
                report["warnings"].append(f"accent page share measured at {accent_share:.1%}; apply a task/profile threshold if needed")

    pdf = None
    if args.pdf:
        if not args.pdf.exists():
            report["failures"].append(f"missing PDF: {args.pdf}")
        else:
            pdf = pdf_metrics(args.pdf, protected_terms_from_manifest(rows or []))
            report["metrics"]["pdf"] = pdf
            if pdf.get("available") is False:
                report["warnings"].append(pdf["warning"])
                report["failures"].append(f"PDF could not be observed: {pdf['warning']}")
            else:
                artifact_pages.append(("PDF", {page["page"]: page for page in pdf["pages"]}))
                if rows is not None and len(rows) != pdf["page_count"]:
                    report["failures"].append(f"manifest rows {len(rows)} != PDF pages {pdf['page_count']}")
                if pdf.get("broken_glyph_pages"):
                    report["failures"].append(f"replacement/missing glyphs on PDF pages {pdf['broken_glyph_pages']}")
                if pdf.get("forbidden_line_breaks"):
                    report["failures"].append(f"forbidden Japanese line breaks: {len(pdf['forbidden_line_breaks'])}")
                if pdf.get("protected_term_splits"):
                    report["failures"].append(f"protected terms split across lines: {len(pdf['protected_term_splits'])}")

    if rows is not None and artifact_pages:
        report["metrics"]["observed"] = {}
        for artifact_label, pages in artifact_pages:
            metrics, failures, warnings = validate_observed_artifact(
                rows,
                pages,
                args.min_observed_visuals,
                args.min_body_pt,
                strict=args.strict,
                artifact_label=artifact_label,
            )
            report["metrics"]["observed"][artifact_label] = metrics
            report["failures"].extend(failures)
            report["warnings"].extend(f"{artifact_label}: {warning}" for warning in warnings)

    report["status"] = "fail" if report["failures"] else "pass"
    payload = json.dumps(report, ensure_ascii=False, indent=2)
    if args.output:
        args.output.write_text(payload + "\n", encoding="utf-8")
    print(payload)
    return 2 if args.strict and report["failures"] else 0


def main():
    """Always turn malformed deck inputs into a structured JSON failure."""
    try:
        return _main_impl()
    except Exception as exc:  # Never expose malformed audit/artifact input as a traceback.
        report = {
            "failures": [{"code": "checker_exception", "message": f"{type(exc).__name__}: {exc}"}],
            "warnings": [],
            "metrics": {},
            "status": "fail",
        }
        print(json.dumps(report, ensure_ascii=False, indent=2, default=str))
        return 2


if __name__ == "__main__":
    sys.exit(main())
